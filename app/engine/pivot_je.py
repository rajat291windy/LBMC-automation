#!/usr/bin/env python3
"""
Revenue JE - Pivot and Summary JE (local version)
=================================================

Does for a Journal Entry export what is done by hand in Excel:

  STEP 1  Read the Journal Entry file (CSV or Excel)
  STEP 2  Fill Debit / Credit from Amount        Amount > 0 -> Debit,  Amount < 0 -> Credit
          (a file that already has DEBIT / CREDIT columns is used as it is)
  STEP 3  Pivot table: Rows = Account, Values = Sum of Debit, Sum of Credit, plus Grand Total
  STEP 4  Summary JE - the table next to the pivot in the Working file:
            11000                          kept as is
            18300                          shown as a second 11000 line
            accounts starting with 1000    (today only 10005) removed - netted into that 11000 line
            13000 and 50005                removed - they cancel; any difference goes into that 11000 line
            all other accounts             kept as is
            Net = Debit - Credit
  STEP 5  Checks: the pivot balances, the summary balances, new accounts
  STEP 6  Save an Excel file next to the input:
            "Pivot and Summary" sheet - the pivot and the Summary JE side by side
                                         (the Summary JE cells are formulas that point at the pivot)
            "Data" sheet              - the file's lines with the Debit / Credit columns filled in

How to run
  In VS Code:  open this folder, set INPUT_PATH below (or leave it empty to try the 7003 sample)
               and press the Run button (the triangle at the top right).
  In a terminal (in the folder that has this file):
    pip install pandas openpyxl
    python pivot_je.py "C:\\Revenue\\7003-Export_JournalEntry.csv"
    python pivot_je.py "C:\\Revenue\\Manual process"                 (every Journal Entry file in a folder)
    python pivot_je.py "C:\\Revenue\\Manual process" --out-dir "C:\\Revenue\\Output"

The original export is only read - it is never changed.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as missing:
    sys.exit(f"The library '{missing.name}' is not installed. In the terminal run:\n"
             f"    python -m pip install pandas openpyxl")

# =============================================================================
# WHICH FILE TO RUN - used when you press Run in VS Code (no path typed)
# =============================================================================
# Paste the full path of your Journal Entry CSV, or of a whole month folder, between the two quotes
# (if you used "Copy as path" in Windows, leave out the quotes it adds). Keep the r in front.
# Leave it empty to run the 7003 sample in the sample_data folder.
#   e.g.  INPUT_PATH = r"C:\Revenue\Export_JournalEntry.csv"
INPUT_PATH = r""

# =============================================================================
# SETTINGS - the business rules (change here if a rule changes)
# =============================================================================
CASH_ACCOUNT = "11000"             # kept as is, and also takes the place of 18300
CLEARING_ACCOUNT = "18300"         # not used in the summary - shown as a second 11000 line
NETTED_PREFIX = "1000"             # accounts starting with this are netted into that 11000 line (today: 10005)
                                   #   -> change to "10" if the whole 10000-series should be netted
CANCEL_PAIR = ("13000", "50005")   # cancel each other out; any difference goes into the 11000 line
KNOWN_ACCOUNTS = {"10005", "11000", "13000", "18300", "21110",
                  "40000", "40005", "40010", "40015", "50005", "65115"}
TOLERANCE = Decimal("0.01")        # allowed difference in the balance checks

ZERO, CENT = Decimal("0.00"), Decimal("0.01")
MONEY = "#,##0.00;(#,##0.00);0.00"  # Excel number format: negatives in brackets, like the Working file
JOURNAL_FILE = re.compile(r"journal\s*_?\s*entry", re.I)
SKIP_FILE = re.compile(r"pivot and summary|working|je summary|^~\$", re.I)


# =============================================================================
# STEP 1 - Read the Journal Entry file
# =============================================================================
def read_file(path: Path) -> pd.DataFrame:
    """Read every column as text, so account numbers and amounts arrive exactly as in the file."""
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                return pd.read_csv(path, dtype=str, keep_default_na=False, encoding=encoding)
            except UnicodeDecodeError:
                continue
        return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="latin-1")
    return pd.read_excel(path, dtype=str, keep_default_na=False)


def find_column(df: pd.DataFrame, *names: str) -> str | None:
    """Find a column by name, ignoring case, spaces and underscores (ACCT_NO = Acct No = acctno)."""
    wanted = {re.sub(r"[^A-Z0-9]", "", n.upper()) for n in names}
    for col in df.columns:
        if re.sub(r"[^A-Z0-9]", "", str(col).upper()) in wanted:
            return col
    return None


def to_money(value) -> Decimal:
    """'1,234.56'  '(57.72)'  '$57.72'  '57.72-'  ''  ->  exact amount with 2 decimals."""
    text = str(value).strip()
    if text == "" or text.lower() in ("nan", "none", "null"):
        return ZERO
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace(",", "").replace("$", "").strip()
    if text.endswith("-"):
        negative, text = True, text[:-1]
    number = Decimal(text)                      # raises InvalidOperation if it is not a number
    return (-number if negative else number).quantize(CENT, rounding=ROUND_HALF_UP)


def clean_account(value) -> str:
    text = str(value).strip()
    return text.split(".")[0] if re.fullmatch(r"\d+\.0+", text) else text   # 18300.0 -> 18300


def account_order(account: str):
    return (0, int(account), "") if account.isdigit() else (1, 0, account)


# =============================================================================
# STEP 2 - Fill Debit / Credit from Amount
# =============================================================================
def fill_debit_credit(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, str, list[str]]:
    """
    Returns
      data     - the file as it is, plus the filled Debit / Credit columns (for the Data sheet)
      lines    - Row, Account, Debit, Credit (used for the pivot)
      fmt      - "1" = Amount only (Debit/Credit filled by the rule), "2" = Debit/Credit given in the file
      problems - anything that needs a look (bad amounts, blank accounts, differences)
    """
    problems: list[str] = []
    account_col = find_column(raw, "Account", "ACCT_NO", "Account Number")
    amount_col = find_column(raw, "Amount")
    debit_col, credit_col = find_column(raw, "Debit"), find_column(raw, "Credit")
    if account_col is None:
        raise ValueError("no Account column found (looked for Account / ACCT_NO)")
    if amount_col is None and (debit_col is None or credit_col is None):
        raise ValueError("no Amount column and no Debit / Credit columns found")

    def parse(column: str) -> list[Decimal | None]:
        out = []
        for i, value in enumerate(raw[column]):
            try:
                out.append(to_money(value))
            except InvalidOperation:
                out.append(None)
                problems.append(f"Row {i + 2}: {column} is not a number ({value!r}) - line left out")
        return out

    data = raw.copy()
    if amount_col is not None:                                   # Format 1 - Amount only
        fmt = "1"
        amounts = parse(amount_col)
        debit = [a if a is not None and a > 0 else ZERO for a in amounts]
        credit = [-a if a is not None and a < 0 else ZERO for a in amounts]
        bad = [a is None for a in amounts]
        new_dr, new_cr = ("Debit", "Credit") if debit_col is None else ("Debit (from Amount)", "Credit (from Amount)")
        data[amount_col] = [a if a is not None else v for a, v in zip(amounts, raw[amount_col])]  # as numbers
        data[new_dr], data[new_cr] = debit, credit
        if debit_col is not None and credit_col is not None:     # the file already had Debit / Credit
            given_dr, given_cr = parse(debit_col), parse(credit_col)
            differ = sum(1 for d, c, gd, gc, b in zip(debit, credit, given_dr, given_cr, bad)
                         if not b and (d != gd or c != gc))
            if differ:
                problems.append(f"{differ} line(s): the Debit/Credit already in the file differ from the Amount")
    else:                                                        # Format 2 - Debit / Credit given
        fmt = "2"
        debit, credit = parse(debit_col), parse(credit_col)
        bad = [d is None or c is None for d, c in zip(debit, credit)]
        data[debit_col] = [d if d is not None else ZERO for d in debit]
        data[credit_col] = [c if c is not None else ZERO for c in credit]

    lines = pd.DataFrame({"Row": range(2, len(raw) + 2),       # Excel row number (row 1 = headings)
                          "Account": [clean_account(a) for a in raw[account_col]],
                          "Debit": debit, "Credit": credit})
    lines = lines[[not b for b in bad]]
    blank = lines["Account"] == ""
    for row in lines.loc[blank & ((lines["Debit"] != ZERO) | (lines["Credit"] != ZERO)), "Row"]:
        problems.append(f"Row {row}: amount with a blank account - line left out")
    return data, lines[~blank].reset_index(drop=True), fmt, problems


# =============================================================================
# STEP 3 - Pivot table: Rows = Account, Values = Sum of Debit, Sum of Credit
# =============================================================================
def build_pivot(lines: pd.DataFrame) -> pd.DataFrame:
    """The same table as the Excel PivotTable (exact to the cent)."""
    pivot = lines.groupby("Account")[["Debit", "Credit"]].sum()
    return pivot.loc[sorted(pivot.index, key=account_order)]


# =============================================================================
# STEP 4 - Summary JE (the adjustments)
# =============================================================================
def build_summary(pivot: pd.DataFrame) -> tuple[list[dict], Decimal, list[str]]:
    """Returns the summary lines, the 13000/50005 difference and the netted accounts."""
    def dr(a): return pivot.at[a, "Debit"] if a in pivot.index else ZERO
    def cr(a): return pivot.at[a, "Credit"] if a in pivot.index else ZERO

    special = {CASH_ACCOUNT, CLEARING_ACCOUNT, *CANCEL_PAIR}
    netted = [a for a in pivot.index if a.startswith(NETTED_PREFIX) and a not in special]
    leftover = sum((dr(a) - cr(a) for a in CANCEL_PAIR), ZERO)    # > 0 extra Debit, < 0 extra Credit
    summary: list[dict] = []

    # 11000 - kept as is
    if CASH_ACCOUNT in pivot.index:
        summary.append({"account": CASH_ACCOUNT, "what": f"{CASH_ACCOUNT} as is",
                        "debit": dr(CASH_ACCOUNT), "credit": cr(CASH_ACCOUNT), "kind": "as_is"})

    # 18300 -> second 11000 line, with the 1000-series netted in and any 13000/50005 difference
    if CLEARING_ACCOUNT in pivot.index or netted or leftover:
        what = f"{CLEARING_ACCOUNT} posted as {CASH_ACCOUNT}"
        if netted:
            what += f", {' + '.join(netted)} netted"
        if leftover:
            what += f", {'/'.join(CANCEL_PAIR)} difference {abs(leftover):,.2f} {'Dr' if leftover > 0 else 'Cr'}"
        summary.append({
            "account": CASH_ACCOUNT, "what": what, "kind": "cash_line",
            "debit": dr(CLEARING_ACCOUNT) + max(leftover, ZERO),
            "credit": (cr(CLEARING_ACCOUNT) - sum((dr(a) for a in netted), ZERO)
                       + sum((cr(a) for a in netted), ZERO) + max(-leftover, ZERO)),
        })

    # every other account - kept as is
    for a in pivot.index:
        if a in special or a in netted:
            continue
        summary.append({"account": a, "debit": dr(a), "credit": cr(a), "kind": "as_is",
                        "what": "kept as is" if a in KNOWN_ACCOUNTS else "NEW account - kept as is, please check"})
    for line in summary:
        line["net"] = line["debit"] - line["credit"]
    return summary, leftover, netted


# =============================================================================
# STEP 5 - Checks
# =============================================================================
def run_checks(pivot, summary, leftover, problems) -> tuple[list[str], str]:
    pd_, pc = sum(pivot["Debit"], ZERO), sum(pivot["Credit"], ZERO)
    sd, sc = sum((s["debit"] for s in summary), ZERO), sum((s["credit"] for s in summary), ZERO)
    new = [a for a in pivot.index if a not in KNOWN_ACCOUNTS]
    results = [
        ("Pivot balances (Grand Total Debit = Credit)", abs(pd_ - pc) <= TOLERANCE, f"difference {pd_ - pc:,.2f}"),
        ("Summary JE balances (Debit = Credit)", abs(sd - sc) <= TOLERANCE, f"difference {sd - sc:,.2f}"),
        (f"{' / '.join(CANCEL_PAIR)} cancel out", True,
         "fully cancelled" if leftover == 0 else f"difference {leftover:,.2f} moved into the {CASH_ACCOUNT} line"),
        ("No new accounts", not new, ", ".join(new) if new else "all accounts known"),
    ]
    cash_line = next((s for s in summary if s["kind"] == "cash_line"), None)
    if cash_line and cash_line["credit"] < 0:
        results.append((f"{CASH_ACCOUNT} line credit not negative", False,
                        "more was deposited than collected - please check"))
    lines = [f"  {'OK    ' if ok else 'CHECK '} {name}: {detail}" for name, ok, detail in results]
    lines += [f"  CHECK  {p}" for p in problems[:20]]
    if len(problems) > 20:
        lines.append(f"  CHECK  ... and {len(problems) - 20} more")
    status = "OK" if all(ok for _, ok, _ in results) and not problems else "REVIEW"
    return lines, status


# =============================================================================
# STEP 6 - Excel output: pivot + Summary JE side by side, and the Data sheet
# =============================================================================
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")   # light blue, like the Excel pivot style
TOTAL_FILL = PatternFill("solid", fgColor="BDD7EE")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="9BC2E6")


def write_excel(path: Path, location: str, source: Path, fmt: str, data: pd.DataFrame,
                pivot: pd.DataFrame, summary: list[dict], netted: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot and Summary"
    ws["A1"] = f"Location {location} - Pivot and Summary JE"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = (f"Source: {source.name}  |  Format {fmt} ({'Amount only' if fmt == '1' else 'Debit/Credit given'})"
                f"  |  {len(data):,} lines  |  created {datetime.now():%d %b %Y %H:%M}")
    ws["A2"].font = Font(italic=True, color="666666")

    # ---- the pivot (columns A-C) ----------------------------------------------------------
    top = 4
    for col, text in zip("ABC", ["Row Labels", "Sum of Debit", "Sum of Credit"]):
        ws[f"{col}{top}"] = text
        ws[f"{col}{top}"].font, ws[f"{col}{top}"].fill = BOLD, HEAD_FILL
    row_of = {}                                           # account -> row of the pivot
    r = top
    for r, account in enumerate(pivot.index, start=top + 1):
        row_of[account] = r
        ws[f"A{r}"] = account
        ws[f"B{r}"] = float(pivot.at[account, "Debit"])
        ws[f"C{r}"] = float(pivot.at[account, "Credit"])
    grand = r + 1
    ws[f"A{grand}"] = "Grand Total"
    ws[f"B{grand}"] = f"=ROUND(SUM(B{top + 1}:B{grand - 1}),2)"
    ws[f"C{grand}"] = f"=ROUND(SUM(C{top + 1}:C{grand - 1}),2)"
    for col in "ABC":
        ws[f"{col}{grand}"].font, ws[f"{col}{grand}"].fill = BOLD, TOTAL_FILL
        ws[f"{col}{grand}"].border = Border(top=THIN)
        for rr in range(top + 1, grand + 1):
            if col != "A":
                ws[f"{col}{rr}"].number_format = MONEY

    # ---- the Summary JE (columns F-J), formulas that point at the pivot -----------------------
    def ref(col: str, account: str) -> str | None:
        return f"{col}{row_of[account]}" if account in row_of else None

    def join(terms: list[tuple[str, str | None]]) -> str:
        parts = [f"{sign}{cell}" for sign, cell in terms if cell]
        return "".join(parts).lstrip("+") or "0"

    pair = [a for a in CANCEL_PAIR if a in row_of]
    leftover_f = "+".join(f"({ref('B', a)}-{ref('C', a)})" for a in pair) or "0"
    for col, text in zip("FGHIJ", ["Account", "Debit", "Credit", "Net", "What this line is"]):
        ws[f"{col}{top}"] = text
        ws[f"{col}{top}"].font, ws[f"{col}{top}"].fill = BOLD, HEAD_FILL
    r = top
    for r, line in enumerate(summary, start=top + 1):
        ws[f"F{r}"] = line["account"]
        if line["kind"] == "cash_line":
            debit = join([("+", ref("B", CLEARING_ACCOUNT))])
            credit = join([("+", ref("C", CLEARING_ACCOUNT))]
                          + [("-", ref("B", a)) for a in netted] + [("+", ref("C", a)) for a in netted])
            ws[f"G{r}"] = f"=ROUND({debit}+MAX({leftover_f},0),2)"
            ws[f"H{r}"] = f"=ROUND({credit}+MAX(-({leftover_f}),0),2)"
            for col in "FGHIJ":
                ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
        else:
            ws[f"G{r}"] = f"={ref('B', line['account'])}"
            ws[f"H{r}"] = f"={ref('C', line['account'])}"
        ws[f"I{r}"] = f"=ROUND(G{r}-H{r},2)"
        ws[f"J{r}"] = line["what"]
        for col in "GHI":
            ws[f"{col}{r}"].number_format = MONEY
    total = r + 1
    ws[f"F{total}"] = "Total"
    ws[f"G{total}"] = f"=ROUND(SUM(G{top + 1}:G{total - 1}),2)"
    ws[f"H{total}"] = f"=ROUND(SUM(H{top + 1}:H{total - 1}),2)"
    ws[f"I{total}"] = f"=ROUND(G{total}-H{total},2)"
    ws[f"J{total}"] = f'=IF(ABS(I{total})<=0.01,"Balanced","NOT BALANCED - check")'
    for col in "FGHIJ":
        ws[f"{col}{total}"].font, ws[f"{col}{total}"].fill = BOLD, TOTAL_FILL
        ws[f"{col}{total}"].border = Border(top=THIN)
    for col in "GHI":
        ws[f"{col}{total}"].number_format = MONEY

    note = total + 2
    ws[f"F{note}"] = "How the Summary JE is built from the pivot:"
    ws[f"F{note}"].font = BOLD
    rules = [f"{CASH_ACCOUNT} kept as is.",
             f"{CLEARING_ACCOUNT} is shown as a second {CASH_ACCOUNT} line. Accounts starting with {NETTED_PREFIX} "
             f"({', '.join(netted) or 'none in this file'}) are netted into it: Credit = {CLEARING_ACCOUNT} Credit "
             f"- their Debit + their Credit.",
             f"{' and '.join(CANCEL_PAIR)} cancel each other out; any difference goes into that {CASH_ACCOUNT} line.",
             "All other accounts are kept as is.  Net = Debit - Credit."]
    for i, text in enumerate(rules, start=1):
        ws[f"F{note + i}"] = f"{i}. {text}"
    for col, width in {"A": 14, "B": 16, "C": 16, "D": 3, "E": 3, "F": 11, "G": 16, "H": 16, "I": 16, "J": 58}.items():
        ws.column_dimensions[col].width = width
    for rr in range(top, total + 1):
        ws[f"A{rr}"].alignment = ws[f"F{rr}"].alignment = Alignment(horizontal="left")
        ws[f"J{rr}"].alignment = Alignment(horizontal="left", indent=1)
    for col in "BCGHI":                                    # amount headings line up with the amounts
        ws[f"{col}{top}"].alignment = Alignment(horizontal="right")
    ws.page_setup.orientation = "landscape"                # prints on one page width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    # ---- the Data sheet: the file's lines with Debit / Credit filled in ------------------------
    wd = wb.create_sheet("Data")
    wd.append([str(c) for c in data.columns])
    money_cols = [i for i, c in enumerate(data.columns, start=1)
                  if re.sub(r"[^A-Z]", "", str(c).upper())
                  in ("AMOUNT", "DEBIT", "CREDIT", "DEBITFROMAMOUNT", "CREDITFROMAMOUNT")]
    for values in data.itertuples(index=False):
        wd.append([float(v) if isinstance(v, Decimal) else v for v in values])
    for i in money_cols:
        letter = wd.cell(row=1, column=i).column_letter
        wd.column_dimensions[letter].width = 14
        for cell in wd[letter][1:]:
            cell.number_format = MONEY
    for cell in wd[1]:
        cell.font, cell.fill = BOLD, HEAD_FILL
    wd.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True                   # Excel works out the formulas when it opens the file
    wb.save(path)


# =============================================================================
# Run one file
# =============================================================================
def money(d: Decimal) -> str:
    return f"({abs(d):,.2f})" if d < 0 else f"{d:,.2f}"


def process_file(path: Path, out_dir: Path | None, used_names: set[str]) -> dict:
    raw = read_file(path)                                                     # STEP 1
    data, lines, fmt, problems = fill_debit_credit(raw)                       # STEP 2
    pivot = build_pivot(lines)                                                # STEP 3
    summary, leftover, netted = build_summary(pivot)                          # STEP 4
    check_lines, status = run_checks(pivot, summary, leftover, problems)      # STEP 5

    loc_col = find_column(raw, "Location", "LOCATION_ID")
    values = [v.strip() for v in raw[loc_col]] if loc_col else []
    values = [v for v in values if v]
    if values:
        location = max(set(values), key=values.count)
    else:
        m = re.match(r"(\d+)", path.stem)
        location = m.group(1) if m else path.stem

    # print the same two tables as in the Working file
    print(f"\nLocation {location}  -  {path.name}  (Format {fmt}, {len(raw):,} lines)")
    print("-" * 78)
    print(f"{'Row Labels':<14}{'Sum of Debit':>18}{'Sum of Credit':>18}")
    for account in pivot.index:
        print(f"{account:<14}{money(pivot.at[account, 'Debit']):>18}{money(pivot.at[account, 'Credit']):>18}")
    pd_, pc = sum(pivot["Debit"], ZERO), sum(pivot["Credit"], ZERO)
    print(f"{'Grand Total':<14}{money(pd_):>18}{money(pc):>18}")
    print()
    print(f"{'Summary JE':<14}{'Debit':>16}{'Credit':>16}{'Net':>16}   What this line is")
    for s in summary:
        print(f"{s['account']:<14}{money(s['debit']):>16}{money(s['credit']):>16}{money(s['net']):>16}   {s['what']}")
    sd, sc = sum((s["debit"] for s in summary), ZERO), sum((s["credit"] for s in summary), ZERO)
    print(f"{'Total':<14}{money(sd):>16}{money(sc):>16}{money(sd - sc):>16}")
    print("\nChecks")
    print("\n".join(check_lines))

    target_dir = out_dir or path.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    name = f"{path.stem} - Pivot and Summary.xlsx"
    if name.lower() in used_names:                           # two inputs with the same name (folder run)
        name = f"{path.stem} ({path.parent.name}) - Pivot and Summary.xlsx"
    used_names.add(name.lower())
    target = target_dir / name
    try:
        write_excel(target, location, path, fmt, data, pivot, summary, netted)   # STEP 6
        print(f"\nSaved: {target}")
    except PermissionError:
        print(f"\nCould not save {target} - is it open in Excel? Close it and run again.")
        status = "REVIEW"
    print(f"Status: {status}")
    return {"location": location, "file": path.name, "format": fmt, "lines": len(raw),
            "pivot": pd_, "summary": sd, "status": status}


def find_files(folder: Path) -> list[Path]:
    return sorted(p for p in folder.rglob("*")
                  if p.is_file() and p.suffix.lower() in (".csv", ".xlsx")
                  and JOURNAL_FILE.search(p.stem) and not SKIP_FILE.search(p.name))


def main() -> int:
    global NETTED_PREFIX
    parser = argparse.ArgumentParser(description="Pivot and Summary JE for Journal Entry exports.")
    parser.add_argument("path", nargs="?",
                        help="a Journal Entry CSV/XLSX file, or a folder (every Journal Entry file inside); "
                             "if left out, INPUT_PATH at the top of this file is used")
    parser.add_argument("--out-dir", help="where to save the Excel files (default: next to each input file)")
    parser.add_argument("--netted-prefix", default=NETTED_PREFIX,
                        help=f"accounts starting with this are netted into the 11000 line (default {NETTED_PREFIX})")
    args = parser.parse_args()
    NETTED_PREFIX = args.netted_prefix

    here = Path(__file__).resolve().parent
    if args.path:                                                 # typed in the terminal
        path = Path(args.path.strip('"')).expanduser()
    elif INPUT_PATH.strip():                                      # set at the top of this file
        path = Path(INPUT_PATH.strip().strip('"')).expanduser()
        if not path.is_absolute():
            path = here / path
    else:                                                         # nothing given: run the sample
        path = here / "sample_data" / "7003-Export_JournalEntry.csv"
        print("No file given - running the 7003 sample. To run your own file, put its path in INPUT_PATH\n"
              "at the top of pivot_je.py, or type:  python pivot_je.py \"<your file or folder>\"")
    out_dir = Path(args.out_dir).expanduser() if args.out_dir else None
    if path.is_dir():
        files = find_files(path)
    elif path.is_file():
        files = [path]
    else:
        print(f"Not found: {path}")
        return 2
    if not files:
        print(f"No Journal Entry files found in {path}")
        return 2

    results, used = [], set()
    for f in files:
        try:
            results.append(process_file(f, out_dir, used))
        except Exception as exc:                                  # one bad file never stops the others
            print(f"\n{f.name}: ERROR - {exc}")
            results.append({"location": "", "file": f.name, "format": "", "lines": 0,
                            "pivot": ZERO, "summary": ZERO, "status": "ERROR"})
    if len(results) > 1:
        print("\n" + "=" * 78)
        print(f"{'Location':<10}{'Format':<8}{'Lines':>7}{'Pivot total':>17}{'Summary total':>17}  Status")
        for r in results:
            print(f"{r['location'] or '-':<10}{r['format'] or '-':<8}{r['lines']:>7,}"
                  f"{money(r['pivot']):>17}{money(r['summary']):>17}  {r['status']:<7} {r['file']}")
        counts = {s: sum(r["status"] == s for r in results) for s in ("OK", "REVIEW", "ERROR")}
        print(f"\n{len(results)} files: {counts['OK']} OK, {counts['REVIEW']} REVIEW, {counts['ERROR']} ERROR")
    return 0 if all(r["status"] == "OK" for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
